#!/usr/bin/env python3
"""
PPTX Report Agent: Fills the Sikich status-report PowerPoint template with
prioritized KPIs.

Role: Deterministic template renderer (no LLM call)
Input: Prioritized JSON from the Prioritization Agent
Output: One-slide branded PPTX status report

Unlike the HTML Status Report Agent, this is intentionally NOT an LLM call.
templates/status_report_template.pptx is the actual Sikich-branded PowerPoint
template (named placeholder shapes, official theme/master), so the safer and
more consistent way to fill it is direct python-pptx text/color substitution
rather than asking a model to freehand a slide layout every time.
"""

import copy
import json
from pathlib import Path
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

TEMPLATE_PATH = Path(__file__).parent.parent / "templates" / "status_report_template.pptx"

STATUS_COLORS = {
    "on_track": RGBColor(0x38, 0x8E, 0x3C),
    "ahead": RGBColor(0x38, 0x8E, 0x3C),
    "at_risk": RGBColor(0xFB, 0xC0, 0x2D),
    "delayed": RGBColor(0xD3, 0x2F, 0x2F),
}
DEFAULT_STATUS_COLOR = RGBColor(0x9E, 0x9E, 0x9E)  # not assessed / unknown

STATUS_LABELS = {
    "on_track": "On Track",
    "ahead": "Ahead of Plan",
    "at_risk": "At Risk",
    "delayed": "Delayed",
    None: "Not Assessed",
}

CONTRACT_TYPE_LABELS = {
    "fixed_fee": "Fixed Fee",
    "time_and_materials": "Time & Materials (T&M)",
}


def _fmt(value: Any, default: str = "TBD") -> str:
    return default if value in (None, "", []) else str(value)


def _set_run_text(paragraph, text: str) -> None:
    """Overwrite a paragraph's visible text while keeping its first run's
    formatting (font size/bold/color) and dropping any extra runs, so
    substituted text inherits the template's styling exactly."""
    runs = paragraph.runs
    if not runs:
        return
    runs[0].text = text
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)


def _set_bullet_list(text_frame, header: str, items: List[str], placeholder_if_empty: str = "None reported") -> None:
    """Keep paragraph[0] as the (bold) header, and replace the template's
    sample bullet paragraphs with one paragraph per item, cloning the
    template's first bullet paragraph's XML so indent/bullet/font formatting
    carries over regardless of how many items are supplied."""
    paragraphs = text_frame.paragraphs
    if paragraphs:
        _set_run_text(paragraphs[0], header)

    if len(paragraphs) < 2:
        return  # template has no bullet paragraph to clone from

    template_p = copy.deepcopy(paragraphs[1]._p)
    parent = paragraphs[1]._p.getparent()

    # Remove all existing bullet paragraphs (everything after the header).
    for p in list(paragraphs[1:]):
        parent.remove(p._p)

    display_items = items if items else [placeholder_if_empty]
    for item_text in display_items:
        new_p = copy.deepcopy(template_p)
        t_elem = new_p.find(f".//{qn('a:t')}")
        if t_elem is not None:
            t_elem.text = item_text
        parent.append(new_p)


def _set_bar_color(shape, status_key: Optional[str]) -> None:
    color = STATUS_COLORS.get(status_key, DEFAULT_STATUS_COLOR)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _find_shape(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise KeyError(f"Template shape not found: {name!r}")


class PptxReportAgent:
    """Fills the Sikich PPTX status-report template from prioritized KPIs."""

    def __init__(self, template_path: Path = TEMPLATE_PATH):
        self.template_path = Path(template_path)

    def generate_pptx_report(self, prioritized_json_path: str, output_pptx_path: str) -> str:
        with open(prioritized_json_path, "r", encoding="utf-8") as f:
            data: Dict[str, Any] = json.load(f)

        prs = Presentation(str(self.template_path))
        slide = prs.slides[0]

        project_name = data.get("source_kpis", {}).get("project_name") or "Project Status Report"
        _set_run_text(_find_shape(slide, "Title 1").text_frame.paragraphs[0], project_name)

        header_line = (
            f"Updated: {_fmt(data.get('status_date'))}  |  "
            f"Go-Live: {_fmt(data.get('go_live_date'))}  |  "
            f"Project Closure: {_fmt(data.get('project_closure_date'))}"
        )
        _set_run_text(_find_shape(slide, "Text Placeholder 10").text_frame.paragraphs[0], header_line)

        status_indicators = data.get("status_indicators") or {}
        for dimension, box_name, bar_name in [
            ("scope", "Scope status", "Scope status bar"),
            ("timeline", "Timeline status", "Timeline status bar"),
            ("resources", "Resources status", "Resources status bar"),
        ]:
            status_value = status_indicators.get(dimension)
            box = _find_shape(slide, box_name)
            _set_run_text(box.text_frame.paragraphs[1], STATUS_LABELS.get(status_value, STATUS_LABELS[None]))
            _set_bar_color(_find_shape(slide, bar_name), status_value)

        metrics = data.get("metrics") or {}
        budget_box = _find_shape(slide, "Budget status")
        budget_status_value = status_indicators.get("budget")
        _set_run_text(budget_box.text_frame.paragraphs[1], STATUS_LABELS.get(budget_status_value, STATUS_LABELS[None]))
        contract_type = CONTRACT_TYPE_LABELS.get(metrics.get("contract_type"), "Contract Type: TBD")
        _set_run_text(budget_box.text_frame.paragraphs[2], contract_type)
        hours_line = (
            f"{_fmt(metrics.get('hours_estimated'))} est | "
            f"{_fmt(metrics.get('hours_consumed'))} consumed | "
            f"{_fmt(metrics.get('hours_remaining'))} remaining"
        )
        _set_run_text(budget_box.text_frame.paragraphs[3], hours_line)
        _set_bar_color(_find_shape(slide, "Budget status bar"), budget_status_value)

        upcoming = data.get("upcoming_focus") or []
        up_next_items = [
            f"{item.get('milestone', 'TBD')} — {_fmt(item.get('due_date'))} ({_fmt(item.get('owner'), 'Unassigned')})"
            for item in upcoming[:4]
        ]
        _set_bullet_list(_find_shape(slide, "Text Placeholder 2").text_frame, "Up Next", up_next_items)

        accomplishments = (data.get("key_accomplishments") or [])[:4]
        _set_bullet_list(_find_shape(slide, "Text Placeholder 3").text_frame, "Value Delivered", accomplishments)

        # Keep each line to a short label (title + severity only, no rationale prose) -
        # this column is narrow and one-slide format has no room for full sentences.
        risks_items = [
            f"{r.get('title', 'Risk')} ({r.get('severity', 'n/a')})"
            for r in (data.get("risks_to_highlight") or [])
        ]
        risks_items += list(data.get("critical_blockers") or [])
        _set_bullet_list(_find_shape(slide, "Text Placeholder 4").text_frame, "Risks / Issues", risks_items[:3])

        stakeholders = data.get("key_stakeholders") or []
        stakeholder_items = [
            f"{s.get('name', 'Unknown')}" + (f" — {s['role']}" if s.get("role") else "")
            for s in stakeholders[:5]
        ]
        _set_bullet_list(_find_shape(slide, "Text Placeholder 5").text_frame, "Key Stakeholders", stakeholder_items)

        output_path = Path(output_pptx_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output_path))

        print("✅ PPTX Report generated!")
        print(f"\U0001f4be Saved to {output_path}")

        return str(output_path)


def main():
    import argparse

    parser = argparse.ArgumentParser(description="Generate PPTX status report from prioritized KPIs")
    parser.add_argument("--input", required=True, help="Path to prioritized_output.json")
    parser.add_argument("--output", required=True, help="Output PPTX file path")
    parser.add_argument("--template", default=str(TEMPLATE_PATH), help="Path to the PPTX template")

    args = parser.parse_args()

    agent = PptxReportAgent(template_path=Path(args.template))
    agent.generate_pptx_report(args.input, args.output)


if __name__ == "__main__":
    main()
